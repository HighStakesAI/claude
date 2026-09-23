#!/usr/bin/env python3
"""Render the fillable form as a .docx.

Used for the Word version clients can be emailed, and as the high-fidelity
fallback if the HTML upload ever imports badly: dragging this into Drive and
opening it with Google Docs reproduces the template exactly.

Palette is the production site's (production-site/style.css).

Run:  python3 make-fillable.py && python3 build-docx.py
In:   Universal-Client-Intake-Form.fillable.md
Out:  Universal-Client-Intake-Form.docx
"""
import re
import sys

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

import mdblocks
import theme

SRC = 'Universal-Client-Intake-Form.fillable.md'
OUT = 'Universal-Client-Intake-Form.docx'

INLINE = re.compile(r'(\*\*.+?\*\*|\*[^*\n]+?\*)')
REQUIRED = re.compile(r'(\[REQUIRED TO START[^\]]*\])')


def _el(tag, **attrs):
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn('w:' + k), v)
    return e


def shade(cell, color):
    cell._tc.get_or_add_tcPr().append(
        _el('w:shd', val='clear', color='auto', fill=color))


def table_borders(table, color, size=6):
    borders = _el('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        borders.append(_el('w:' + edge, val='single', sz=str(size),
                           space='0', color=color))
    table._tbl.tblPr.append(borders)


def row_height(row, inches):
    row._tr.get_or_add_trPr().append(
        _el('w:trHeight', val=str(int(inches * 1440)), hRule='atLeast'))


def cell_margins(table):
    mar = _el('w:tblCellMar')
    for side, val in (('top', 60), ('bottom', 60), ('left', 110), ('right', 110)):
        mar.append(_el('w:' + side, w=str(val), type='dxa'))
    table._tbl.tblPr.append(mar)


def para_border(paragraph, edge, color, size=6, space=4):
    pPr = paragraph._p.get_or_add_pPr()
    bdr = pPr.find(qn('w:pBdr'))
    if bdr is None:
        bdr = _el('w:pBdr')
        pPr.append(bdr)
    bdr.append(_el('w:' + edge, val='single', sz=str(size),
                   space=str(space), color=color))


def keep_with_next(paragraph):
    paragraph._p.get_or_add_pPr().append(_el('w:keepNext', val='true'))


def add_runs(paragraph, text, bold=False, italic=False, color=None, size=None):
    """Inline Markdown, with [REQUIRED TO START] picked out in the brand red."""
    for chunk in INLINE.split(text):
        if not chunk:
            continue
        b, i = bold, italic
        if chunk.startswith('**') and chunk.endswith('**'):
            chunk, b = chunk[2:-2], True
        elif chunk.startswith('*') and chunk.endswith('*') and len(chunk) > 2:
            chunk, i = chunk[1:-1], True
        for part in REQUIRED.split(chunk):
            if not part:
                continue
            run = paragraph.add_run(part)
            run.bold = b
            run.italic = i
            if REQUIRED.fullmatch(part):
                run.bold = True
                run.font.color.rgb = RGBColor.from_string(theme.RED.upper())
            elif color:
                run.font.color.rgb = RGBColor.from_string(color.upper())
            if size:
                run.font.size = Pt(size)


def build(md_path=SRC, out_path=OUT):
    doc = Document()
    for section in doc.sections:
        section.top_margin = section.bottom_margin = Inches(0.8)
        section.left_margin = section.right_margin = Inches(0.85)

    normal = doc.styles['Normal']
    normal.font.name = 'Arial'
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(theme.INK.upper())
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    normal.paragraph_format.line_spacing = 1.15

    stats = {'boxes': 0, 'tables': 0, 'headings': 0}

    for block in mdblocks.parse(open(md_path, encoding='utf-8').read()):
        kind = block[0]

        if kind == 'h1':
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
            add_runs(p, block[1], bold=True, size=21)
            stats['headings'] += 1

        elif kind == 'h2':
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(20)
            p.paragraph_format.space_after = Pt(9)
            add_runs(p, block[1], bold=True, size=13)
            para_border(p, 'bottom', theme.RED, size=10, space=3)
            keep_with_next(p)
            stats['headings'] += 1

        elif kind == 'rule':
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            para_border(p, 'bottom', theme.BORDER, size=6)

        elif kind == 'quote':
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.16)
            p.paragraph_format.space_before = Pt(2)
            add_runs(p, block[1], italic=True, color=theme.MUTED)
            para_border(p, 'left', theme.BORDER, size=12, space=7)

        elif kind == 'bullets':
            for item in block[1]:
                p = doc.add_paragraph(style='List Bullet')
                p.paragraph_format.space_after = Pt(3)
                p.paragraph_format.left_indent = Inches(0.3)
                add_runs(p, item)

        elif kind == 'box':
            table = doc.add_table(rows=1, cols=1)
            table_borders(table, theme.BORDER)
            cell_margins(table)
            cell = table.rows[0].cells[0]
            shade(cell, theme.BOX_FILL)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
            cell.paragraphs[0].paragraph_format.space_after = Pt(0)
            row_height(table.rows[0],
                       theme.BOX_TALL_IN if block[1] else theme.BOX_IN)
            doc.add_paragraph().paragraph_format.space_after = Pt(4)
            stats['boxes'] += 1

        elif kind == 'table':
            rows = block[1]
            table = doc.add_table(rows=len(rows), cols=len(rows[0]))
            table_borders(table, theme.BORDER)
            cell_margins(table)
            for r, row in enumerate(rows):
                if r:
                    row_height(table.rows[r], theme.BOX_IN)
                for c, text in enumerate(row):
                    if c >= len(table.columns):
                        continue
                    cell = table.rows[r].cells[c]
                    cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_after = Pt(0)
                    add_runs(p, text, bold=(r == 0))
                    shade(cell, theme.HEAD_FILL if r == 0 else theme.BOX_FILL)
            doc.add_paragraph().paragraph_format.space_after = Pt(4)
            stats['tables'] += 1

        elif kind == 'para':
            text = block[1]
            p = doc.add_paragraph()
            add_runs(p, text)
            if text.startswith('**'):
                keep_with_next(p)
                p.paragraph_format.space_before = Pt(8)
                p.paragraph_format.space_after = Pt(4)

    doc.save(out_path)
    print(f"wrote {out_path}: {stats['boxes']} shaded answer boxes, "
          f"{stats['tables']} data tables, {stats['headings']} headings")
    return stats


if __name__ == '__main__':
    try:
        build()
    except FileNotFoundError:
        sys.exit(f'{SRC} not found — run make-fillable.py first.')
